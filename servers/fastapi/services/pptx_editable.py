
"""R1 F13: editable PPTX from slide UI (text/tables/charts), not slide screenshots."""

from __future__ import annotations

import os
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from utils.asset_directory_utils import get_exports_directory

EDITOR_W = 1280
EDITOR_H = 720
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "stacked_bar": XL_CHART_TYPE.COLUMN_STACKED,
    "donut": XL_CHART_TYPE.DOUGHNUT,
    "waterfall": XL_CHART_TYPE.COLUMN_CLUSTERED,
}


def _iter_elements(tree: Any):
    if isinstance(tree, dict):
        if tree.get("type"):
            yield tree
        for value in tree.values():
            yield from _iter_elements(value)
    elif isinstance(tree, list):
        for item in tree:
            yield from _iter_elements(item)


def _text_from_element(element: dict[str, Any]) -> str:
    chunks: list[str] = []
    for run in element.get("runs") or []:
        if isinstance(run, dict) and run.get("text"):
            chunks.append(str(run["text"]))
        elif isinstance(run, str):
            chunks.append(run)
    if element.get("text"):
        chunks.append(str(element["text"]))
    return "".join(chunks).strip()


def _box(element: dict[str, Any]) -> tuple:
    pos = element.get("position") or {}
    size = element.get("size") or {}
    x = float(pos.get("x") or 0) / EDITOR_W * SLIDE_W
    y = float(pos.get("y") or 0) / EDITOR_H * SLIDE_H
    w = max(float(size.get("width") or 200) / EDITOR_W * SLIDE_W, Inches(0.4))
    h = max(float(size.get("height") or 40) / EDITOR_H * SLIDE_H, Inches(0.3))
    return x, y, w, h


def _add_text(slide, element: dict[str, Any]) -> None:
    text = _text_from_element(element)
    if not text:
        return
    x, y, w, h = _box(element)
    box = slide.shapes.add_textbox(x, y, w, h)
    box.text_frame.word_wrap = True
    box.text_frame.text = text


def _add_table(slide, element: dict[str, Any]) -> None:
    rows = element.get("rows") or []
    cols = element.get("columns") or []
    if not rows and not cols:
        return
    def cell_text(cell) -> str:
        if isinstance(cell, dict):
            return _text_from_element(cell) or str(cell.get("text") or "")
        return str(cell or "")
    header = [cell_text(c) for c in cols] if cols else []
    body = [[cell_text(c) for c in row] for row in rows] if rows and isinstance(rows[0], list) else []
    n_cols = max([len(header)] + [len(r) for r in body] or [1])
    n_rows = (1 if header else 0) + len(body)
    if n_rows == 0:
        return
    x, y, w, h = _box(element)
    table = slide.shapes.add_table(n_rows, n_cols, x, y, w, h).table
    r_i = 0
    if header:
        for c_i, value in enumerate(header + [""] * n_cols):
            if c_i >= n_cols:
                break
            table.cell(0, c_i).text = value
        r_i = 1
    for row in body:
        for c_i in range(n_cols):
            table.cell(r_i, c_i).text = row[c_i] if c_i < len(row) else ""
        r_i += 1


def _waterfall_stacks(values: list[float]) -> tuple[list[float], list[float], list[str]]:
    base: list[float] = []
    visible: list[float] = []
    colors: list[str] = []
    running = 0.0
    for value in values:
        if value >= 0:
            base.append(running)
            visible.append(value)
            colors.append("up")
            running += value
        else:
            running += value
            base.append(running)
            visible.append(-value)
            colors.append("down")
    return base, visible, colors


def _add_waterfall(slide, element: dict[str, Any]) -> None:
    categories = [str(c) for c in (element.get("categories") or [])]
    series = element.get("series") or []
    if not categories or not series:
        return
    values = []
    for raw in (series[0].get("values") if isinstance(series[0], dict) else []) or []:
        try:
            values.append(float(raw))
        except (TypeError, ValueError):
            values.append(0.0)
    base, visible, colors = _waterfall_stacks(values)
    x, y, w, h = _box(element)
    n = max(1, len(visible))
    gap = max(1, int(w * 0.06 / n))
    bar_w = max(8, int((w - gap * (n + 1)) / n))
    peak = max((b + v) for b, v in zip(base, visible)) or 1.0
    labels = categories[:n] or [f"S{i+1}" for i in range(n)]
    for i, (b, v, c, label) in enumerate(zip(base, visible, colors, labels)):
        bx = x + gap + i * (bar_w + gap)
        bh = max(8, int(h * 0.82 * (v / peak)))
        by = y + int(h * 0.82 * (1 - (b + v) / peak))
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, by, bar_w, bh)
        shape.fill.solid()
        shape.fill.fore_color.rgb = (
            RGBColor(0x12, 0xB7, 0x6A) if c == "up" else RGBColor(0xEF, 0x44, 0x44)
        )
        try:
            shape.line.fill.background()
        except Exception:
            pass
        cap = slide.shapes.add_textbox(bx, y + int(h * 0.86), bar_w, int(h * 0.12))
        cap.text_frame.text = str(label)[:16]


def _add_chart(slide, element: dict[str, Any]) -> None:
    categories = [str(c) for c in (element.get("categories") or [])]
    series = element.get("series") or []
    if not categories or not series:
        return
    if str(element.get("chart_type") or "") == "waterfall":
        _add_waterfall(slide, element)
        return
    chart_type = CHART_TYPES.get(str(element.get("chart_type") or "bar"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    data = CategoryChartData()
    data.categories = categories
    for item in series:
        if not isinstance(item, dict):
            continue
        values = []
        for raw in item.get("values") or []:
            try:
                values.append(float(raw))
            except (TypeError, ValueError):
                values.append(0.0)
        data.add_series(str(item.get("name") or "Series"), values)
    x, y, w, h = _box(element)
    slide.shapes.add_chart(chart_type, x, y, w, h, data)


def build_editable_pptx(*, title: str, slides: list[dict[str, Any]], dest_path: str) -> str:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    for slide_data in slides:
        slide = prs.slides.add_slide(blank)
        ui = slide_data.get("ui") or {}
        found = False
        for element in _iter_elements(ui):
            kind = element.get("type")
            if kind == "text":
                _add_text(slide, element)
                found = True
            elif kind == "table":
                _add_table(slide, element)
                found = True
            elif kind == "chart":
                _add_chart(slide, element)
                found = True
            elif kind == "infographic":
                labels = []
                for node in element.get("nodes") or []:
                    if isinstance(node, dict):
                        labels.append(str(node.get("label") or node.get("id") or ""))
                if not labels:
                    labels = [str(x) for x in (element.get("items") or []) if x]
                caption = "Infographic (placeholder): " + (", ".join(x for x in labels if x) or "unlabelled")
                box = slide.shapes.add_textbox(Inches(0.4), Inches(6.8), Inches(12), Inches(0.4))
                box.text_frame.text = caption[:200]
                found = True
        note = (slide_data.get("speaker_note") or "").strip()
        if note and not found:
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(6))
            box.text_frame.text = note
        elif not found:
            title_text = str((slide_data.get("content") or {}).get("title") or "")
            if title_text:
                box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
                box.text_frame.text = title_text
    os.makedirs(os.path.dirname(dest_path) or ".", exist_ok=True)
    prs.save(dest_path)
    return dest_path
